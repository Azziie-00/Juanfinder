from pathlib import Path
from io import BytesIO
import fitz
from docx import Document

MAX_FILE_BYTES = 10 * 1024 * 1024
MAX_TEXT_CHARS = 60_000
ALLOWED_EXTENSIONS = {".pdf", ".docx", ".txt"}

def extract_document(filename: str, content: bytes) -> str:
    suffix = Path(filename).suffix.lower()
    if suffix not in ALLOWED_EXTENSIONS:
        raise ValueError("Unsupported file type. Upload PDF, DOCX, or TXT.")
    if not content:
        raise ValueError("The uploaded file is empty.")
    if len(content) > MAX_FILE_BYTES:
        raise ValueError("File exceeds the 10 MB limit.")

    if suffix == ".pdf":
        chunks = []
        with fitz.open(stream=content, filetype="pdf") as pdf:
            if len(pdf) > 30:
                raise ValueError("PDFs may contain at most 30 pages.")
            for page in pdf:
                chunks.append(page.get_text("text"))
        text = "\n".join(chunks).strip()
        if len(text) < 40:
            raise ValueError(
                "This PDF appears to be scanned or contains too little selectable text. "
                "Please upload a text-based PDF or DOCX/TXT version. OCR can be enabled separately."
            )
    elif suffix == ".docx":
        doc = Document(BytesIO(content))
        chunks = [p.text for p in doc.paragraphs if p.text.strip()]
        for table in doc.tables:
            for row in table.rows:
                chunks.append(" | ".join(cell.text.strip() for cell in row.cells))
        text = "\n".join(chunks).strip()
    else:
        text = content.decode("utf-8-sig", errors="replace").strip()

    if not text:
        raise ValueError("No readable text was found in this file.")
    return text[:MAX_TEXT_CHARS]
